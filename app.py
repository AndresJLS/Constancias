import os  #os es para crear carpetas y manejar archivos
from docxtpl import DocxTemplate #aqui se importa la libreria para manejar plantillas de docx
from openpyxl import load_workbook

CARPETA_DOCX = "temp_docx"
CARPETA_PDF = "constancias_pdf"
os.makedirs(CARPETA_DOCX, exist_ok=True) #esto crea la carpeta temp_docx si no existe, exist_ok=True evita que de error si ya existe
os.makedirs(CARPETA_PDF, exist_ok=True) #lo mismo que lo anterior pero para las constancias

hoja = load_workbook("equipos.xlsx").active

# Paso A: un .docx por equipo, nombrado con el número de equipo
for fila in hoja.iter_rows(min_row=2, values_only=True):
    equipo, integrantes, tema, lugar, fecha = fila
    if not equipo:
        continue
    equipo = str(equipo).strip() #esto es para quitar espacios en blanco al inicio y al final del número de equipo

    plantilla = DocxTemplate("Plantilla.docx")
    plantilla.render({"equipo": equipo, "integrantes": integrantes, "tema": tema, "lugar": lugar, "fecha": fecha}) #aqui se renderiza la plantilla con los datos del equipo
    plantilla.save(f"{CARPETA_DOCX}/Constancia_Equipo_{equipo}.docx") #aqui se guarda el archivo .docx con el nombre del equipo

# Paso B: convertir TODOS los .docx a PDF en una sola llamada
from docx2pdf import convert
convert(CARPETA_DOCX, CARPETA_PDF)