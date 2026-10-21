import os
import sys
import tempfile
from datetime import date, datetime
if sys.stdout is None:
    sys.stdout = open(os.devnull, "w")
if sys.stderr is None:
    sys.stderr = open(os.devnull, "w")

from docx2pdf import convert
from docxtpl import DocxTemplate
from openpyxl import load_workbook
from PyQt5.QtWidgets import (QApplication, QFileDialog, QLabel, QMessageBox,
                             QPushButton, QWidget)

CARPETA_PDF = "constancias_pdf"


def texto_fecha(valor):
    """Si Excel entrega la fecha como fecha real, la pasamos a texto legible."""
    if isinstance(valor, (datetime, date)): #aqui se verifica si el valor es una fecha, si es asi se formatea a dd/mm/yyyy
        return valor.strftime("%d/%m/%Y")#luego se regresa el valor formateado
    return valor


def generar_constancias(ruta_excel):
    """Lee el Excel, crea un PDF por equipo y regresa cuántos creó."""
    os.makedirs(CARPETA_PDF, exist_ok=True)
    hoja = load_workbook(ruta_excel).active      # <- abrir el libro; antes faltaba esto

    with tempfile.TemporaryDirectory() as carpeta_docx:   # los .docx son solo intermedios, este paso los borra al terminar
        cantidad = 0
        # Paso A: un .docx por equipo, nombrado con el número de equipo
        for fila in hoja.iter_rows(min_row=2, values_only=True):
            equipo, integrantes, tema, lugar, fecha = fila[:5]   # [:5] ignora columnas de más
            if not equipo:#si no hay equipo, se salta la fila
                continue
            equipo = str(equipo).strip()

            plantilla = DocxTemplate("Plantilla.docx")
            plantilla.render({"equipo": equipo, "integrantes": integrantes, "tema": tema,
                              "lugar": lugar, "fecha": texto_fecha(fecha)})#importamos la funcion texto_fecha para que si la fecha es un objeto de fecha, se formatee a dd/mm/yyyy
            plantilla.save(os.path.join(carpeta_docx, f"Constancia_Equipo_{equipo}.docx"))#aqui se guarda el archivo .docx con el nombre del equipo en la carpeta temporal
            cantidad += 1 #se pone un contador para saber cuántos equipos se procesaron

        # Paso B: convertir TODOS los .docx a PDF en una sola llamada
        if cantidad:
            convert(carpeta_docx, CARPETA_PDF)
    return cantidad


if __name__ == "__main__":
    app = QApplication(sys.argv)
    w = QWidget()
    w.resize(300, 300)
    w.setWindowTitle("Constancias de equipos")

    label = QLabel("Selecciona el archivo de equipos (.xlsx)", w)
    label.move(40, 120)

    btn2 = QPushButton("Seleccionar archivo", w)
    btn2.move(90, 160)

    def seleccionar_archivo():
        archivo, _ = QFileDialog.getOpenFileName(
            w, "Seleccionar archivo", "", "Archivos de Excel (*.xlsx)")
        if not archivo:
            return                        # si cancelan, simplemente no pasa nada
        try:
            cantidad = generar_constancias(archivo)   # aquí se usa la ruta
        except Exception as e:
            QMessageBox.critical(w, "Error", f"{type(e).__name__}: {e}")
            return
        QMessageBox.information(w, "Listo", f"Se generaron {cantidad} constancias en '{CARPETA_PDF}'.")

    btn2.clicked.connect(seleccionar_archivo)

    w.show()
    sys.exit(app.exec_())